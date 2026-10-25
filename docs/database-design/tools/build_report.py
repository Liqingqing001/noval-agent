from pathlib import Path
import json
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TABLES = json.loads((ROOT / 'schema-metadata.json').read_text(encoding='utf-8'))
RESULT = json.loads((ROOT / 'validation-results.json').read_text(encoding='utf-8'))
FONT = 'Microsoft YaHei'
FONTPATH = 'C:/Windows/Fonts/msyh.ttc'
OUT = ROOT / 'NovelAgent数据库设计报告_李晴_202400702062.docx'

def draw_diagram():
    im = Image.new('RGB', (2400, 1550), 'white')
    d = ImageDraw.Draw(im)
    ft = ImageFont.truetype(FONTPATH, 38)
    fs = ImageFont.truetype(FONTPATH, 27)
    small = ImageFont.truetype(FONTPATH, 26)
    boxes = {}
    def box(key, x, y, label, sub, width=270):
        boxes[key] = (x, y, x+width, y+140)
        d.rounded_rectangle(boxes[key], radius=13, fill='#F2F5F8', outline='#314455', width=3)
        d.text((x+width/2,y+33),label,font=ft,fill='#000000',anchor='mm')
        d.text((x+width/2,y+80),key,font=fs,fill='#223344',anchor='mm')
        d.text((x+width/2,y+115),sub,font=small,fill='#3C4C59',anchor='mm')
    def edge(parent,child,label='1 对 N'):
        a,b=boxes[parent],boxes[child]
        start=((a[0]+a[2])/2,a[3]); end=((b[0]+b[2])/2,b[1])
        mid=(start[1]+end[1])/2
        d.line([start,(start[0],mid),(end[0],mid),end],fill='#536779',width=3)
        d.polygon([(end[0],end[1]),(end[0]-8,end[1]-16),(end[0]+8,end[1]-16)],fill='#536779')
        d.text((end[0]+12,end[1]-42),label,font=small,fill='#1B374D')
    box('app_user',1050,50,'用户账户','PK id')
    box('novel',680,340,'小说','PK id  FK user_id')
    box('model_profile',1910,340,'模型配置','PK id  FK user_id')
    box('model_credential',1910,550,'模型凭据','PK FK profile_id')
    children=[('chapter','章节正文'),('story_character','角色'),('world_setting','世界观'),('volume_outline','卷大纲'),('foreshadowing','伏笔'),('chat_session','对话会话'),('novel_reading_state','阅读状态')]
    xs=[50,390,730,1070,1410,1750,2090]
    for (key,label),x in zip(children,xs): box(key,x,970,label,'FK novel_id')
    for key,label,x,sub in [('chapter_revision','历史版本',50,'FK chapter_id'),('character_alias','角色别名',390,'复合主键'),('chapter_outline','章节大纲',1070,'FK volume_id 可空'),('chat_message','对话消息',1750,'FK session_id')]: box(key,x,1340,label,sub)
    edge('app_user','novel'); edge('app_user','model_profile'); edge('model_profile','model_credential','1 对 0 或 1')
    for key,_ in children: edge('novel',key,'1 对 0 或 1' if key=='novel_reading_state' else '1 对 N')
    edge('chapter','chapter_revision');edge('story_character','character_alias');edge('volume_outline','chapter_outline');edge('chat_session','chat_message')
    # Main containment paths are displayed here; cross references have a dedicated cardinality table.
    im.save(ROOT/'assets'/'er-overview.png')

def flow_diagram():
    im=Image.new('RGB',(1800,530),'white');d=ImageDraw.Draw(im)
    ft=ImageFont.truetype(FONTPATH,30);small=ImageFont.truetype(FONTPATH,23)
    labels=[('浏览器编辑','本地草稿与版本号'),('后端 API','验证身份与拥有者'),('事务保存函数','版本比较与行锁'),('PostgreSQL','正文和版本快照')]
    for i,(a,b) in enumerate(labels):
        x=35+i*450;d.rounded_rectangle((x,150,x+370,340),radius=12,outline='#314455',fill='#F2F5F8',width=3)
        d.text((x+185,205),a,font=ft,fill='black',anchor='mm');d.text((x+185,275),b,font=small,fill='#34495E',anchor='mm')
        if i<3:
            d.line((x+370,245,x+430,245),fill='#536779',width=4);d.polygon([(x+430,245),(x+415,235),(x+415,255)],fill='#536779')
    d.text((900,420),'成功返回新版本号  冲突返回 409 并保留用户草稿',font=ft,fill='#111111',anchor='mm')
    im.save(ROOT/'assets'/'save-flow.png')

draw_diagram();flow_diagram()
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.3)
for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[style_name];style.font.name=FONT;style.font.color.rgb=RGBColor(0,0,0)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(4)
doc.styles['Subtitle'].font.italic=False
for name,size in [('Title',24),('Subtitle',13),('Heading 1',17),('Heading 2',12.5),('Heading 3',11)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(5)
    doc.styles[name].paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0];header.text='深圳技术大学    数据库课程设计    NovelAgent';header.runs[0].font.size=Pt(8)
header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('李晴  202400702062    第 ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' 页').font.size=Pt(8)
doc.core_properties.title='NovelAgent小说创作系统数据库设计报告'
doc.core_properties.author='李晴'
doc.core_properties.subject='数据库课程大作业 需求分析 概念设计 关系模式 规范化 物理设计 SQL 验证'

pending_break=False
def p(text='',style=None):
    global pending_break
    par=doc.add_paragraph(text,style)
    if pending_break: par.paragraph_format.page_break_before=True;pending_break=False
    return par
def h(text,level=1):
    global pending_break
    par=doc.add_heading(text,level)
    if pending_break: par.paragraph_format.page_break_before=True;pending_break=False
    return par
def page():
    global pending_break
    pending_break=True
def section(landscape=False):
    global pending_break
    pending_break=False
    s=doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width=Inches(11 if landscape else 8.5);s.page_height=Inches(8.5 if landscape else 11)
    s.top_margin=Inches(.62 if landscape else .72);s.bottom_margin=Inches(.62 if landscape else .72)
    s.left_margin=Inches(.7);s.right_margin=Inches(.7)
def code(text):
    for line in text.strip('\n').split('\n'):
        par=p();par.paragraph_format.space_after=Pt(1);par.paragraph_format.line_spacing=1.0
        r=par.add_run(line);r.font.name='Consolas';r.font.size=Pt(8.2)
        r._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
def table(headers,rows,widths=None,size=9):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [7.1/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for i,value in enumerate(headers): t.rows[0].cells[i].text=str(value)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row): cells[i].text=str(value)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if ri==0: repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'25' if side in ['top','bottom'] else '85');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF4' if ri==0 else ('F8FAFC' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0);par.paragraph_format.space_before=Pt(0);par.paragraph_format.line_spacing=1.05
                for r in par.runs:
                    r.font.name=FONT;r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0);r.bold=(ri==0)
    p().paragraph_format.space_after=Pt(1)
    return t
def bullet(text): p('• '+text)
def picture(name,width=7.05): doc.add_picture(str(ROOT/'assets'/name),width=Inches(width))

# Cover
p('数据库课程大作业', 'Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('\n')
title=p('NovelAgent小说创作系统\n数据库设计报告','Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('面向多用户上线的 PostgreSQL 数据库设计与 SQL 验证','Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('\n')
for value in ['学校  深圳技术大学','专业  数据科学与大数据技术','年级  2024级','姓名  李晴','学号  202400702062','日期  2026年10月25日']:
    par=p(value);par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.space_after=Pt(12)
p('\n')
p('成果类型  课程设计与隔离 SQL 验证稿').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('项目已有 SQLite 实现  本报告给出 PostgreSQL 目标方案').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('未修改现有应用数据库  未声明已完成生产部署').alignment=WD_ALIGN_PARAGRAPH.CENTER
page()
h('摘要与阅读说明')
p('本设计以已有的 NovelAgent 小说创作项目为业务案例，规划面向多用户部署的关系数据库。系统围绕用户、小说、章节、创作资料、AI 对话与模型配置组织数据，并通过章节版本记录降低误覆盖风险。方案沿用 Express 与 Prisma 的应用结构，目标数据库采用 PostgreSQL 16 及以上版本。')
p('设计形成 15 张业务表和 133 个字段，提供实体关系图、关系模式、函数依赖分析、数据字典、索引与约束、事务保存函数，以及可执行 SQL。SQL 在独立 PGlite 内存 PostgreSQL 中完成 45 项验证，全部通过；这不替代独立服务器上的并发、性能与备份恢复测试。')
p('关键词  关系数据库  概念设计  第三范式  参照完整性  乐观锁  PostgreSQL')
h('报告结构',2)
for text in ['一 项目背景与需求分析','二 概念结构设计','三 逻辑结构与规范化','四 物理设计与访问路径','五 数据字典','六 完整性与事务处理','七 SQL 查询与操作示例','八 验证结果与测试边界','九 部署迁移与运行维护','十 设计总结与答辩准备','参考资料与交付清单']:
    p(text)
p('阅读顺序建议  先理解业务规则和 ER 图，再对照数据字典阅读 SQL。图中给出主干联系，所有可选关联和同小说约束以关系表及建表脚本为准。')
page()
h('一 项目背景与需求分析')
h('1 项目现状与设计目标',2)
p('现有应用采用 React 与 TypeScript 前端、Express 后端以及 Prisma ORM。后端已经建立用户、小说、章节、角色、世界观、大纲和伏笔等 SQLite 表。小说等实体具备后端接口，AI 对话与模型配置仍主要存放于浏览器 localStorage。')
p('本课程设计的目标是把这些分散状态转化为清晰的持久化模型，说明每个实体的边界，建立能够校验关联归属、记录正文历史并支持跨设备恢复的数据基础。目标方案属于后续实施规划，报告中的新表并未自动接入现有应用。')
table(['现状','存在的问题','设计应对'],[
 ['角色别名存为 JSON 字符串','不便逐个去重与查询','拆为 character_alias'],
 ['卷纲与章纲共用 PlotCard','类型和父子关系混在一起','拆为 volume_outline 与 chapter_outline'],
 ['多个关联 ID 只是字符串','可能引用其他小说的实体','增加同小说复合外键'],
 ['对话仅保存在浏览器','换设备后无法恢复历史','增加会话与消息表'],
 ['正文没有独立版本关系','误覆盖后缺少恢复依据','增加章节快照与版本保存事务'],
 ['模型密钥与配置放在一起','同步时容易暴露凭据','配置和加密凭据分表'],
],[1.65,2.45,3.0])
h('2 用户角色和业务范围',2)
p('第一版采用单作者拥有一本小说的模型：一位用户可以拥有多本小说，一本小说只有一位拥有者。用户可以维护章节、角色、世界观和大纲，追踪伏笔，建立 AI 对话并选择自己的模型配置。运维人员负责备份、迁移和监控，不通过普通业务页面读取用户模型密钥。')
p('本期不设计多人实时协作、付费订阅、公共作品交易、向量检索和模型计费结算。若后续支持共同作者，应增加小说成员及权限关系，不能直接允许任意用户修改 novel.user_id。')
page()
h('3 功能需求与关键业务规则',2)
table(['编号','需求或规则','主要落点'],[
 ['R01','邮箱归一化后唯一 密码只保存哈希','app_user'],
 ['R02','小说及其内容只能由拥有者访问','novel.user_id 与后端身份校验'],
 ['R03','章节同书顺序唯一 支持事务换序','chapter 延迟唯一约束'],
 ['R04','章纲可未归卷 可暂未关联正文','可空 volume_id 与 chapter_id'],
 ['R05','角色和伏笔的章节锚点必须属于同书','复合外键'],
 ['R06','完成的章纲必须关联正文','CHECK'],
 ['R07','回收的伏笔必须有回收位置','CHECK'],
 ['R08','用户可保存多个模型 仅一个默认配置','部分唯一索引'],
 ['R09','会话不能引用其他用户的模型配置','消息归属检查触发器'],
 ['R10','重复生成请求不能重复写入消息','request_id 唯一约束'],
 ['R11','正文保存必须检测客户端版本','save_chapter 与版本触发器'],
 ['R12','历史恢复必须形成新版本','chapter_revision 与 restore 来源'],
],[.55,4.35,2.2])
h('4 非功能需求与假设',2)
p('数据正确性优先于写入成功提示。保存失败应保留浏览器草稿；服务器确认提交后才显示云端保存成功。所有对外 API 都需要执行身份、拥有者和输入白名单校验，外键不能替代访问权限校验。')
p('教学容量假设为 1000 位用户、每人平均 3 本小说、每本 100 章、每章约 5000 个字符，总计约 15 亿字符。若以每个中文字符约 3 字节粗估，正文约 4.5 GB；此数值不含索引、对话、行存储开销和历史快照，也不代表压缩后的实际磁盘占用。版本保留策略会显著影响容量。')
p('所有时间使用带时区时间戳；业务序号从 0 开始；ID 使用 UUID。字数定义为应用统一处理后的正文字符数，需去除 HTML 标记和约定的空白字符，不等于模型 Token 数。')
section(True)
h('二 概念结构设计')
picture('er-overview.png',9.4)
p('图 1 总体 ER 图  共 15 个实体  附加交叉联系见下一页')
section(False)
h('实体和联系的识别',2)
p('实体按创作内容、AI 对话和账户配置组织。ER 图展示主干联系，英文实体名与数据字典对应。跨小说校验所需的附加外键集中在下表说明。')
p('图中 N 表示多侧可以有零个或多个实例。chapter_outline 可以未归卷，此时仍直接属于 novel；chat_message 可以不引用模型配置，例如用户消息。消息中引用配置的权限规则不在图中用交叉线展开，以避免遮挡主干联系。')
h('联系与基数',2)
table(['联系','基数','约束或解释'],[
 ['用户到小说','1 对 N','拥有者必填 一书一主'],['用户到模型配置','1 对 N','配置名称在同一用户内唯一'],
 ['小说到章节和资料','1 对 N','子实体必须归属一本小说'],['小说到阅读状态','1 对 0或1','拥有者由小说推导'],
 ['角色到别名','1 对 N','同角色别名不重复'],['卷纲到章纲','1 对 N','章纲允许未归卷'],
 ['正文到章纲','1 对 0或1','章纲的 chapter_id 唯一 可为空'],['正文到历史版本','1 对 N','创建即保存第一个快照'],
 ['小说到会话','1 对 N','会话不得跨小说移动'],['会话到消息','1 对 N','sequence 在会话内唯一'],
 ['模型配置到凭据','1 对 0或1','免认证本地模型无需凭据行'],['模型配置到消息','1 对 N','删除配置后引用置空 快照保留'],
 ['正文或章纲到伏笔','1 对 N','分别表示埋设和回收 可选锚点'],
],[2.0,.9,4.2])
h('同小说联系的表达',2)
p('仅校验 chapter_id 存在，并不能证明章节属于当前小说。因此章节表提供 UNIQUE(novel_id, id)，引用方以 FOREIGN KEY(novel_id, chapter_id) 建立联系。卷纲、章纲和伏笔采用同样方法。该设计把“存在且归属正确”交给数据库检查。')
p('可选锚点为空时，复合外键使用默认 MATCH SIMPLE，不要求存在目标；novel_id 本身始终非空且受小说外键约束。删除单个被引用章节时，需要在同一事务中先解绑这些可选锚点再删除。删除整本小说时，其内容和状态作为聚合统一级联清理。')
page()
h('三 逻辑结构与规范化')
h('1 关系模式',2)
p('下列为精简逻辑模式。PK 为主键，FK 为外键，UQ 为候选键或唯一性约束；全部非键属性见数据字典。')
for s in [
 'app_user(id PK, email UQ, password_hash, nickname, status, …)',
 'novel(id PK, user_id FK, title, genre, summary, status, …)',
 'chapter(id PK, novel_id FK, title, content, sequence, version, …)',
 'story_character(id PK, novel_id FK, name, role, first_chapter_id FK, …)',
 'character_alias(character_id PK FK, alias PK)',
 'world_setting(id PK, novel_id FK, category, title, content, …)',
 'volume_outline(id PK, novel_id FK, title, summary, sequence, …)',
 'chapter_outline(id PK, novel_id FK, volume_id FK, chapter_id FK UQ, …)',
 'foreshadowing(id PK, novel_id FK, title, status, 四个可选锚点 FK, …)',
 'model_profile(id PK, user_id FK, name, protocol, model_id, 参数, …)',
 'model_credential(profile_id PK FK, ciphertext, nonce, key_version, …)',
 'chat_session(id PK, novel_id FK, title, …)',
 'chat_message(id PK, session_id FK, role, content, sequence, request_id UQ, …)',
 'chapter_revision(id PK, chapter_id FK, version, 正文快照, …)',
 'novel_reading_state(novel_id PK FK, last_chapter_id FK, cursor_offset, …)',
]:
    par=p(s);par.paragraph_format.space_after=Pt(5)
h('2 关键候选键',2)
table(['关系','候选键或业务唯一键'],[
 ['app_user','id 和归一化 email'],['character_alias','character_id 与 alias 的组合'],
 ['chapter','id 和 novel_id 与 sequence 的组合'],['volume_outline','id 和 novel_id 与 sequence 的组合'],
 ['model_profile','id 和 user_id 与 name 的组合'],['chapter_revision','id 和 chapter_id 与 version 的组合'],
 ['chat_message','id 和 session_id 与 sequence 的组合 request_id 为可空唯一幂等标识'],
],[1.8,5.3])
page()
h('3 函数依赖与范式分析',2)
p('逻辑设计以第三范式为目标。第一范式要求属性值在本系统业务语义中不可再分，第二范式消除非主属性对候选键的部分依赖，第三范式要求每个非平凡函数依赖 X → A 中，X 为超键或 A 为主属性。不能因为每张表都有 id，就直接断言全部满足第三范式。')
table(['分析对象','主要函数依赖','设计结论'],[
 ['用户','id → email password_hash 等 email → id','归一化邮箱唯一 避免重复用户事实'],
 ['角色别名','character_id alias 共同标识一行','无其他非键属性 不保留别名列表列'],
 ['章节','id → 当前正文 版本 顺序','不复制小说标题或拥有者姓名'],
 ['模型凭据','profile_id → 密文 随机数 密钥版本','与普通模型配置分离 减少访问面'],
 ['历史快照','chapter_id version → 当时正文与来源','每个版本作为独立历史事实'],
],[1.25,2.8,3.05])
h('4 分解示例与可控冗余',2)
p('若把 user_email、novel_title、chapter_title 和 content 放在一张写作总表中，会出现 chapter_id → novel_id → novel_title 的传递依赖。同一小说改名需要更新多章，可能产生更新异常。将用户、小说、章节分解后，非键事实仅在所属实体保存；章节通过外键连接小说，查询时再连接获取标题。')
p('角色别名原来若以逗号拼接，无法稳定按别名去重和建立外键。拆分为 character_alias(character_id, alias) 后，一个别名占一行，以复合主键约束重复。JSON 数组虽然可作为一个值存储，但本业务把别名作为需单独检索的集合，因此选择拆表，而不是仅以存储格式判断范式。')
p('物理模型为落实同小说复合外键，在章节大纲等关系保留 novel_id，同时引用卷纲或正文。对于已关联的行，volume_id 或 chapter_id 又可推导 novel_id，存在受约束的重复归属事实；因此本报告不声称物理模式的每个关系均严格达到第三范式。这是用少量冗余换取声明式完整性校验的明确取舍，所属关系不可随意移动。')
p('word_count 是为列表展示保存的应用侧派生统计，必须与正文在同一保存事务更新；历史模型名和协议是事件发生时快照，不应与当前配置强制相等。extra_body 与 extra_headers 是供应商扩展文档，不用于存放重复的核心关系或明文凭据。')
page()
h('四 物理设计与访问路径')
h('1 类型选择',2)
table(['类型或结构','使用位置','选择理由'],[
 ['UUID','实体标识和外键','便于客户端预分配并在迁移中保留标识'],['TEXT','正文 简介 提示词','不人为截断长文本'],
 ['VARCHAR','标题 邮箱 状态标识','限制业务输入长度 配合 CHECK'],['TIMESTAMPTZ','创建 更新 完成时间','存储时间点 显示时转换用户时区'],
 ['BIGINT','正文版本号','单调递增作为并发校验依据'],['JSONB','供应商扩展参数','兼容不同接口的小规模结构化扩展'],
 ['BYTEA','凭据密文与随机数','只存二进制密文 不存加密主密钥'],
],[1.2,2.2,3.7])
h('2 索引规划',2)
table(['访问需求','索引或约束','注意事项'],[
 ['用户按最近更新列小说','novel(user_id, updated_at DESC)','避免全表扫描所有用户作品'],
 ['按序读取章节','UQ novel_id sequence','复用唯一索引 不重复建立'],
 ['按角色类别展示','story_character(novel_id, role, sequence)','筛选列在前 排序列在后'],
 ['待写章纲 未回收伏笔','novel_id 与 status 组合索引','需结合真实查询选择性评估'],
 ['最近会话列表','chat_session(novel_id, updated_at DESC)','使用 ID 作为稳定排序补充'],
 ['会话消息游标分页','UQ session_id sequence','用 sequence > 游标 避免深分页'],
 ['默认模型配置','user_id 上的部分唯一索引','仅 is_default=true 的行参与'],
 ['章节历史版本','UQ chapter_id version','支持定位版本和逆序读取'],
],[1.55,2.45,3.1])
p('主键和唯一约束已自动建立必要索引。外键列是否加索引取决于连接、删除及查询需求，不能机械地对每一列建索引。本期不建立大正文索引；全文检索需明确中文分词和查询方式后再扩展。性能验证应使用有代表性数据和 EXPLAIN ANALYZE，本报告不虚构响应时间或并发容量。')
page()
h('五 数据字典')
p('字典由生成 SQL 的同一份字段元数据构建。所有约束以 sql/01_schema.sql 为准；为便于排版，此处用短写展示空值、键和默认值。时间默认“当前时刻”，UUID 默认“自动生成”。')

def dictionary(name):
    t=next(t for t in TABLES if t['name']==name)
    h(t['label']+' '+t['name'],2)
    if t['note']: p(t['note']+'。')
    rows=[]
    for n,typ,label,rules in t['fields']:
        flags=[]
        if 'PRIMARY KEY' in rules: flags.append('PK')
        if 'REFERENCES' in rules: flags.append('FK')
        flags.append('必填' if 'NOT NULL' in rules or 'PRIMARY KEY' in rules else '可空')
        if 'DEFAULT' in rules:
            default=rules.split('DEFAULT',1)[1].split(' CHECK')[0].strip()
            default=default.replace('gen_random_uuid()','自动UUID').replace('CURRENT_TIMESTAMP','当前时刻').replace("'{}'::jsonb",'空对象').replace("''",'空串')
            flags.append('默认'+default)
        rows.append([n,typ,' '.join(flags),label])
    table(['字段','类型','空值及默认值','含义'],rows,[2.0,1.1,1.85,2.15],9)
    if t['constraints']:
        short=[]
        for c in t['constraints']:
            if c.startswith('FOREIGN KEY'): short.append('复合外键 '+c.split(' REFERENCES')[0].replace('FOREIGN KEY ',''))
            elif c.startswith('UNIQUE'): short.append(c.replace(' DEFERRABLE INITIALLY IMMEDIATE',''))
            elif c.startswith('PRIMARY KEY'): short.append(c)
            else: short.append('业务 CHECK 见完整 SQL')
        par=p('表级约束  '+'；'.join(short));par.paragraph_format.space_after=Pt(4)
        for run in par.runs:run.font.size=Pt(8.5)

groups=[['app_user','novel'],['chapter','chapter_revision'],['story_character','character_alias','world_setting'],['volume_outline','chapter_outline'],['foreshadowing','novel_reading_state'],['model_profile'],['model_credential','chat_session'],['chat_message']]
for index,group in enumerate(groups):
    if index>0: page()
    for name in group:dictionary(name)
    if group==['model_profile']:
        p('is_default 使用部分唯一索引实现每用户至多一个默认配置，不要求用户必须存在默认配置。切换默认值时，在一个事务中先取消旧默认，再设置新默认。')
    if group==['chat_message']:
        p('request_id 对每个逻辑生成请求保持不变，重复请求收到唯一约束冲突后，应由服务返回已有结果，而不是重复调用模型。该约束只防止重复写库，不能自行保证上游模型恰好调用一次。')
        p('模型用量未知时保持 NULL，不能把未知值当作零。streaming 消息可以保存阶段正文，但只有成功终态才允许进入自动结构化同步；失败或取消不得冒充完整结果。')
page()
h('六 完整性与事务处理')
h('1 完整性层次',2)
table(['层次','实现','示例'],[
 ['实体完整性','主键 唯一键 非空','用户ID 邮箱 会话内顺序'],['参照完整性','普通及复合外键','章纲不能指向其他小说的正文'],
 ['域完整性','数据类型 CHECK','非负字数 温度范围 协议枚举'],['跨表业务规则','触发器 后端事务','消息的配置属于小说拥有者'],
 ['访问控制','已验证身份和 API 查询过滤','以请求身份决定 user_id'],['生成幂等','请求ID唯一与状态处理','重复请求返回已存在消息'],
],[1.05,2.5,3.55])
p('普通 CHECK 只校验当前行，不能用它代替跨表归属检查。本设计使用复合外键表达同小说关联，用触发器检查跨用户模型引用。拥有者与子实体所属关系建立后禁止直接移动，避免关联行随后失去归属一致性。[1]')
h('2 正文保存事务',2)
p('客户端提交章节 ID、预期版本号及新正文。后端从已验证 JWT 获取用户 ID，再调用 save_chapter。函数锁定目标章节行、检查拥有者和预期版本，更新正文并递增版本。AFTER 触发器把新正文写入历史表。任一步失败则整个事务回滚。')
p('如果设备甲和乙均持有版本 1，甲先保存得到版本 2；乙再以版本 1 请求时，数据库返回 SQLSTATE 40001。后端应区分本函数的过期编辑冲突并转换为 HTTP 409，保留乙的草稿并提示合并。不能自动读取新版本后覆盖重试，否则会把版本冲突重新变成静默丢稿。')
p('普通读写采用 PostgreSQL 默认 Read Committed 隔离级别，并以行锁和版本条件控制正文竞争。对真正的数据库可重试序列化失败，可在无外部副作用且重新读取状态的边界内有限重试；多设备编辑冲突应交给用户处理。[2]')
h('3 排序和删除',2)
p('章节顺序唯一约束声明为可延迟。在同一事务中 SET CONSTRAINTS ALL DEFERRED，再交换两章 sequence，提交时统一检查最终顺序。对于拖拽整批排序，推荐按同样方式一次批量提交。')
p('删除单章需要先清除人物初登场、章纲正文关联、阅读位置和伏笔锚点；若章纲完成状态依赖正文，解绑时同时改回 pending。删除整本小说则级联清理小说内聚合。用户账户删除使用 RESTRICT，先执行明确的数据清理流程；上线可增加软删除和恢复期。')
page()
h('七 SQL 查询与操作示例')
p('完整初始化见 01_schema.sql，示例数据见 02_seed.sql，10 个查询见 03_queries.sql。示例邮箱均使用保留的 example.test 域名，密码仅为不可登录的占位文本，凭据字节也不是实际加密结果。')
h('1 连接查询与聚合',2)
code('''SELECT n.title, s.chapter_count, s.total_words
FROM novel_agent.novel n
JOIN novel_agent.novel_word_stats s ON s.novel_id=n.id
WHERE n.user_id=$1
ORDER BY n.updated_at DESC, n.id;

SELECT n.title, COUNT(c.id), SUM(c.word_count)
FROM novel_agent.novel n
JOIN novel_agent.chapter c ON c.novel_id=n.id
WHERE n.user_id=$1
GROUP BY n.id,n.title HAVING COUNT(c.id)>=2;''')
p('第一个查询复用汇总视图，章节为空时返回零；第二个查询演示连接、分组和 HAVING。真实 API 使用驱动参数绑定，不把用户输入拼接进 SQL。')
h('2 子查询和窗口函数',2)
code('''SELECT n.id,n.title FROM novel_agent.novel n
WHERE n.user_id=$1 AND NOT EXISTS (
  SELECT 1 FROM novel_agent.chapter c WHERE c.novel_id=n.id
);

SELECT title,word_count,
  SUM(word_count) OVER (
    PARTITION BY novel_id ORDER BY sequence
    ROWS UNBOUNDED PRECEDING
  ) AS cumulative_words
FROM novel_agent.chapter WHERE novel_id=$1
ORDER BY sequence;''')
p('NOT EXISTS 检索没有章节的小说，窗口函数保留每章行并计算累计字数。连接多个一对多关系后再直接求和可能导致重复计数，因此总字数先在单独视图中按章节聚合。')
page()
h('3 正文保存与历史恢复',2)
code('''SELECT novel_agent.save_chapter(
  $1,  -- 后端验证后的用户 UUID
  $2,  -- 章节 UUID
  $3,  -- 客户端预期版本
  $4,  -- 新标题
  $5,  -- 新正文
  $6,  -- 按统一口径计算的字数
  $7,  -- html 或 text
  'manual'
);''')
p('函数返回新的版本号。用户编辑器应在返回后更新本地版本，而不是只根据定时器是否结束判断保存成功。首次新建正文时触发器自动保存版本 1。字数由后端采用与编辑器一致的口径计算，数据库检查非负及事务一致更新，不声称可独立验证 HTML 字数。')
code('''SELECT novel_agent.save_chapter(
  $1, $2, $3,
  r.title,r.content,r.word_count,r.content_format,'restore'
)
FROM novel_agent.chapter_revision r
WHERE r.chapter_id=$2 AND r.version=$4;''')
p('恢复旧版本使用其正文作为新输入，以当前预期版本保存。例如从当前版本 2 恢复版本 1 的内容，结果生成版本 3，版本 1 和 2 仍被保留。若目标快照不存在，该查询返回零行，API 应明确提示未找到。')
h('4 模型默认配置与调用状态',2)
code('''BEGIN;
UPDATE novel_agent.model_profile SET is_default=false
WHERE user_id=$1 AND is_default=true;
UPDATE novel_agent.model_profile SET is_default=true
WHERE user_id=$1 AND id=$2;
COMMIT;''')
p('后端需检查第二条 UPDATE 确实影响一行。模型生成是耗时的外部请求，不应在持有正文数据库行锁的长事务里等待模型。先写 queued 消息并提交，再调用服务，最后用短事务更新终态。进程崩溃留下的 queued 或 streaming 状态需要超时任务处理。')
page()
h('八 验证结果与测试边界')
p(f'执行结果为 {RESULT["passed"]} 项通过，{RESULT["failed"]} 项失败。测试引擎为 PGlite 0.5.11 内嵌 PostgreSQL 18.6，内存运行；目标脚本采用 PostgreSQL 16 及以上语法。本次没有启动独立 PostgreSQL 16 服务器，也没有变更项目当前 SQLite 数据。')
table(['测试组','数量','验证内容','结果'],[
 ['T01至T06','6','建表 示例数据 自动快照 汇总与零值','通过'],['T07至T23','17','唯一性 域 外键 归属 幂等 状态约束','通过'],
 ['T24至T30','7','版本保存 快照 过期请求 越权 恢复','通过'],['T31至T32','2','延迟换序与事务回滚','通过'],
 ['T33至T35','3','受引用删除 配置解绑 小说级联','通过'],['Q01至Q10','10','连接 分组 子查询 窗口 游标分页等','通过'],
],[1.25,.5,4.65,.7])
h('关键反例',2)
table(['输入或操作','预期 SQLSTATE','观察结果'],[
 ['重复邮箱或重复请求ID','23505 唯一冲突','被拒绝'],['引用另一小说的章节','23503 外键冲突','被拒绝'],
 ['引用另一用户的模型配置','23514 业务约束','被拒绝'],['过期版本覆盖当前正文','40001 版本冲突','被拒绝'],
 ['其他用户保存章节','42501 权限不足','被拒绝'],['正文变化但版本未增加','23514 版本规则','被拒绝'],
],[3.7,2.0,1.4])
p('逐项结果保存在 validation-results.json，验证程序可从空内存库重新执行。版本冲突通过先后两次使用不同预期版本重放验证，不能据此声称已测试真实多连接并发。SQL 查询测试证明语法可执行，不代表性能已达到某项指标。')
p('尚需在独立服务器验证：多连接保存与死锁重试、迁移大批数据、备份文件可恢复性、实际网络中断、密钥加密和轮换、应用层身份及 CORS、真实模型接口调用。这些是上线验收任务，不作为本次已完成的实验结果。')
page()
h('九 部署迁移与运行维护')
h('1 从现有 SQLite 迁移',2)
table(['阶段','主要工作','验收依据'],[
 ['准备','备份 SQLite 导出实体及关联 建字段映射','原库可恢复 不丢现有文稿'],
 ['建立目标库','建立独立开发 测试 生产数据库','连接凭据与权限彼此隔离'],
 ['新迁移历史','归档旧 SQLite SQL 生成 PostgreSQL 迁移','空库可由迁移完整重建'],
 ['数据转换','旧字符串ID映射UUID 别名拆行 卷章纲拆表','映射表完整 关联无悬空'],
 ['补齐浏览器数据','用户确认后上传对话和模型配置','按用户归属幂等导入'],
 ['校验与切换','比对行数 逐章内容 哈希和关联','抽样及自动校验均通过'],
 ['回退与收尾','保留回退窗口 记录切换点和新写入','不能直接丢弃切换后的新数据'],
],[1.1,3.2,2.8])
p('现有项目主键为 Prisma cuid 或本地字符串，目标为 UUID，因此不能简单改列类型后原样导入。应为每类实体建立 old_id → new_uuid 映射，并在所有子表使用同一映射回填。原 passwordHash 原样保留；不能把已有哈希再次当密码计算哈希。')
p('Prisma 的迁移 SQL 与数据库提供者相关，SQLite 迁移不能直接用于 PostgreSQL。[4] 本报告的独立 SQL 用于课程实验与建模核对；真正接入应用时，应将表、关联和字段映射到 Prisma schema，并把 CHECK、触发器、视图等补入受版本管理的迁移。不能同时让两套未经对齐的建表流程管理同一个 schema。')
h('2 发布流程',2)
p('先在开发环境生成并检查迁移，再在测试环境验证，最后发布生产。项目当前 Prisma 5 使用 migrate deploy 应用已审核迁移；生产环境不运行 migrate dev、reset 或开发数据初始化。前端只调用后端 API，不向浏览器暴露数据库连接串。')
p('按先扩展再切换再清理的方式处理不兼容字段。备份和应用版本都应记录对应的迁移版本；回滚应用并不自动回滚数据结构，删除列等不可逆变更必须单独规划。')
page()
h('3 安全和权限',2)
p('所有查询以已验证的用户身份过滤所属小说，不能使用请求体中的 user_id 决定权限。保存函数的 p_user 是可信后端传入的身份参数，不是用户可以自行声明的数据库登录凭据；因此函数不是替代 JWT 的认证方案。')
p('上线采用独立迁移账号和最小权限运行账号。普通运行账号不应有 CREATE、DROP 或任意 DDL 权限。密码使用应用层自适应密码哈希；模型凭据采用服务端认证加密，主密钥放在环境密钥管理中，模型密钥与密文均不写普通访问日志。')
p('extra_headers 的数据库约束拒绝常见认证字段，应用还须对所有头名称做大小写归一化和白名单处理，禁止通过变体绕过检查。base_url 的 SQL 约束只是基础校验；若后续增加服务端模型代理，还必须独立处理目标地址、重定向、内网访问和超时限制。')
p('可把 PostgreSQL 行级安全作为后续防御，但本报告未启用 RLS，也未验证其策略。使用时需要明确连接池中的身份上下文和表拥有者绕过规则，不能仅打开开关就宣称实现用户隔离。')
h('4 备份与保留',2)
p('建议目标为生产数据库每日自动备份，并根据业务可接受的数据损失配置更细的 WAL 归档或托管服务恢复能力。备份应存放在与数据库故障域分离的位置，定期执行恢复演练。示例目标 RPO 24 小时、RTO 4 小时仅用于课程规划，不是本次实测指标。[3]')
p('章节快照按用户确认的策略保留，例如近期保存较密、较早仅保留每日版本，并为手动标记版本保留例外。快照与主表位于同一个库，不能代替异地备份；用户导出的小说文件也不能替代完整数据库恢复。封面和附件放对象存储，并单独备份元数据与访问策略。')
h('5 分阶段实施',2)
p('第一阶段落地 PostgreSQL、关联约束与可靠保存；第二阶段接入对话和模型配置云同步；第三阶段完善章节恢复、密钥管理、备份和监控。每个阶段先运行隔离测试，再接入现有前端，不在课程设计阶段直接重置线上或当前开发数据。')
page()
h('十 设计总结与答辩准备')
p('本设计将现有小说创作业务转化为可解释的实体与联系，并围绕跨小说引用、重复生成、保存冲突和历史恢复建立数据库规则。课程重点是从需求推导结构、说明范式与冗余的取舍，再用正例和反例验证约束，而不是仅列出建表语句。')
table(['可能提问','回答要点'],[
 ['为什么不用一张大表','用户 小说 章节存在不同函数依赖 合表会造成重复与更新异常'],
 ['有主键是否就满足第三范式','不满足 必须检查所有候选键和非平凡函数依赖'],
 ['为何保留 novel_id 复合外键','落实同小说归属约束 这是明确记录的物理冗余取舍'],
 ['外键能否解决用户越权','只能保证关联合法 仍需认证和拥有者过滤'],
 ['乐观锁和行锁是否重复','版本号识别过期客户端 行锁协调服务器上同时修改'],
 ['保存历史是否等于备份','历史支持业务恢复 同库故障仍会丢失 需要独立备份'],
 ['幂等ID能否保证模型只调用一次','只防止重复数据库记录 外部调用仍需状态与恢复策略'],
 ['测试是否证明已经能上线','没有 仅证明已列 SQL 行为 真实并发与运维还待验收'],
],[2.35,4.75])
page()
h('参考资料与交付清单')
p('[1] PostgreSQL 16 文档 Constraints。用于主键、唯一、CHECK、外键和延迟约束的设计核对。')
p('https://www.postgresql.org/docs/16/ddl-constraints.html')
p('[2] PostgreSQL 16 文档 Transaction Isolation。用于隔离级别和并发冲突处理。')
p('https://www.postgresql.org/docs/16/transaction-iso.html')
p('[3] PostgreSQL 16 文档 Backup and Restore。用于备份及恢复规划。')
p('https://www.postgresql.org/docs/16/backup.html')
p('[4] Prisma 文档 Limitations and known issues。用于跨数据库迁移历史的限制说明。')
p('https://www.prisma.io/docs/orm/v6/prisma-migrate/understanding-prisma-migrate/limitations-and-known-issues')
p('资料访问日期均为 2026年10月25日。项目依据为本仓库 Prisma schema、前端领域类型与持久化逻辑；新增模型均为本报告规划。')
p('交付包括本 Word 报告、3 个 SQL 文件、ER 图与保存流程图、字段元数据、45 项验证记录和可复用生成工具。README 提供运行顺序与验证边界。')

# Disable Word's East Asian document grid for predictable table spacing, and remove theme title rules.
for paragraph in doc.element.iter(qn('w:p')):
    ppr=paragraph.find(qn('w:pPr'))
    if ppr is None: ppr=OxmlElement('w:pPr');paragraph.insert(0,ppr)
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');ppr.append(snap)
for root in [doc.element,doc.styles.element]:
    for border in list(root.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.save(OUT)
print(OUT)
print(f'{len(TABLES)} tables, {sum(len(t["fields"]) for t in TABLES)} dictionary fields')
