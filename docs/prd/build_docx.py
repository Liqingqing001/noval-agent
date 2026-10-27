"""Build the review DOCX from the adjacent Markdown (pandoc + python-docx)."""
from pathlib import Path
import subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
TARGET = ROOT / 'NovelAgent参考作品分析与辅助创作PRD_V1.1.docx'
subprocess.run(['pandoc', str(ROOT / 'content-analysis-creation-prd.md'),
                '--from=gfm', '--to=docx', '-o', str(TARGET)], check=True)
doc = Document(TARGET)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2)

def font(rpr, east='宋体'):
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for key in list(fonts.attrib):
        if 'Theme' in key: del fonts.attrib[key]
    for key, value in [('ascii', 'Calibri'), ('hAnsi', 'Calibri'), ('eastAsia', east), ('cs', 'Calibri')]:
        fonts.set(qn('w:' + key), value)

for style in doc.styles:
    if style.type in (1, 2):
        font(style.element.get_or_add_rPr(), '黑体' if style.name.startswith('Heading') else '宋体')
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.space_after = Pt(6)
for i, size in [(1, 20), (2, 15), (3, 12)]:
    st = next(s for s in doc.styles if s.style_id == f'Heading{i}')
    st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string('173D32')
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
for p in doc.paragraphs:
    for r in p.runs: font(r._r.get_or_add_rPr(), '黑体' if p.style.name.startswith('Heading') else '宋体')
for table in doc.tables:
    table.autofit = False
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'CDD8D1')]:
            border.set(qn('w:' + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    n = len(table.columns)
    ratios = {2: [0.27, 0.73], 3: [0.23, 0.36, 0.41], 4: [0.13, 0.37, 0.17, 0.33]}.get(n, [1/n]*n)
    for column, ratio in zip(table.columns, ratios): column.width = Cm(17 * ratio)
    for ri, row in enumerate(table.rows):
        for cell, ratio in zip(row.cells, ratios):
            cell.width = Cm(17 * ratio)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.keep_with_next = False
                for r in p.runs:
                    font(r._r.get_or_add_rPr()); r.font.size = Pt(9)
                    if ri == 0: r.bold = True
            if ri == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF1ED'); cell._tc.get_or_add_tcPr().append(shade)
        trpr = row._tr.get_or_add_trPr()
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
header = sec.header.paragraphs[0]
header.text = 'NovelAgent  /  产品需求文档  /  V1.1 待评审'
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('第 ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.add_run(' 页')
for p in (header, footer):
    for r in p.runs:
        font(r._r.get_or_add_rPr()); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('64736B')
doc.core_properties.title = 'NovelAgent 参考作品分析与辅助创作 PRD'
doc.core_properties.subject = '书籍拆解、文章分析、文风卡与参考创作'
doc.core_properties.author = 'NovelAgent'
for part in [doc.part, sec.header.part, sec.footer.part]:
    for run in part.element.iter(qn('w:r')):
        rpr = run.get_or_add_rPr()
        if rpr.find(qn('w:rFonts')) is None:
            font(rpr)
doc.save(TARGET)
print(TARGET)
