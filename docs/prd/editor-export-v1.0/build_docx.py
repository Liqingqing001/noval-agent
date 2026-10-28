"""Build the adjacent PRD for Windows Word. Run with: py build_docx.py."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from zipfile import ZipFile
from lxml import etree

ROOT = Path(__file__).resolve().parent
TARGET = ROOT / 'NovelAgent小说编辑与导出PRD_V1.0.docx'


def set_fonts(rpr, east='宋体'):
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attribute in list(fonts.attrib):
        if 'theme' in attribute.lower():
            del fonts.attrib[attribute]
    for key, value in {'ascii': 'Times New Roman', 'hAnsi': 'Times New Roman', 'eastAsia': east, 'cs': 'Times New Roman'}.items():
        fonts.set(qn('w:' + key), value)


def add_text(paragraph, text, east='宋体', bold=False):
    run = paragraph.add_run(text.replace('`', ''))
    run.bold = bold
    set_fonts(run._r.get_or_add_rPr(), east)
    return run


def build():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for style in document.styles:
        if style.type in (1, 2):
            set_fonts(style.element.get_or_add_rPr(), '黑体' if style.name.startswith('Heading') or style.name == 'Title' else '宋体')
    for fonts in list(document.styles.element.iter(qn('w:rFonts'))):
        set_fonts(fonts.getparent(), fonts.get(qn('w:eastAsia')) or '宋体')
    for border in list(document.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 19), ('Heading 1', 14), ('Heading 2', 12)]:
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    lines = (ROOT / 'requirements.md').read_text(encoding='utf-8').splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line.startswith('|'):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip())
                index += 1
            rows = [row for row in rows if not re.fullmatch(r'[| :\-]+', row)]
            table = document.add_table(rows=0, cols=2)
            table.style = 'Table Grid'
            table.autofit = False
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, width, value in zip(cells, (4.5, 12.1), row.strip('|').split('|')):
                    cell.width = Cm(width)
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.25
                    paragraph.paragraph_format.keep_with_next = False
                    run = add_text(paragraph, value.strip(), bold=row_index == 0)
                    run.font.size = Pt(10.5)
                    if row_index == 0:
                        shade = OxmlElement('w:shd')
                        shade.set(qn('w:fill'), 'EDF3EF')
                        cell._tc.get_or_add_tcPr().append(shade)
                row_properties = table.rows[-1]._tr.get_or_add_trPr()
                row_properties.append(OxmlElement('w:cantSplit'))
                if row_index == 0:
                    row_properties.append(OxmlElement('w:tblHeader'))
            continue
        if line.startswith('# '):
            add_text(document.add_paragraph(style='Title'), line[2:], '黑体', True)
        elif line.startswith('## '):
            add_text(document.add_paragraph(style='Heading 1'), line[3:], '黑体', True)
        elif line.startswith('- '):
            add_text(document.add_paragraph(style='List Bullet'), line[2:])
        else:
            add_text(document.add_paragraph(), line)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    add_text(footer, '第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    field_run = OxmlElement('w:r')
    field_text = OxmlElement('w:t')
    field_text.text = '1'
    field_run.append(field_text)
    field.append(field_run)
    footer._p.append(field)
    add_text(footer, ' 页')
    for run in footer.runs:
        set_fonts(run._r.get_or_add_rPr())
        run.font.size = Pt(9)
    for part in (document.part, section.header.part, section.footer.part):
        for run in part.element.iter(qn('w:r')):
            rpr = run.get_or_add_rPr()
            if rpr.find(qn('w:rFonts')) is None:
                set_fonts(rpr)
    document.core_properties.title = 'NovelAgent 小说编辑与导出产品需求'
    document.core_properties.author = '李晴'
    document.core_properties.subject = '段落排版 编辑连续性与小说导出'
    document.save(TARGET)
    # Word may retain a legacy stylesWithEffects part; clear its theme font references too.
    temporary = TARGET.with_suffix('.tmp.docx')
    with ZipFile(TARGET) as source, ZipFile(temporary, 'w') as output:
        for member in source.infolist():
            payload = source.read(member.filename)
            if member.filename.startswith('word/') and member.filename.endswith('.xml'):
                xml = etree.fromstring(payload)
                changed = False
                for fonts in xml.iter(qn('w:rFonts')):
                    for attribute in list(fonts.attrib):
                        if 'theme' in attribute.lower():
                            del fonts.attrib[attribute]
                            changed = True
                    for key, fallback in {'ascii': 'Times New Roman', 'hAnsi': 'Times New Roman', 'eastAsia': '宋体', 'cs': 'Times New Roman'}.items():
                        if fonts.get(qn('w:' + key)) is None:
                            fonts.set(qn('w:' + key), fallback)
                            changed = True
                if changed:
                    payload = etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
            output.writestr(member, payload)
    temporary.replace(TARGET)
    print(TARGET)


if __name__ == '__main__':
    build()
